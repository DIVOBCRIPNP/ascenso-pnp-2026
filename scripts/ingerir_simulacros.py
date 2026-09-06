"""
Parsea Examenes_15x100_PNP_2026.docx → data/simulacros.json

Estructura de salida:
  {
    version: 1,
    total_examenes: 15,
    total_preguntas: 1500,
    examenes: [
      {
        n: 1,
        titulo: "EXAMEN DE SIMULACRO N° 01 · 100 preguntas",
        preguntas: [
          {n: 1, materia: "CONSTITUCIÓN...", pregunta: "...", opciones: [...], correcta: 0}
        ]
      }, ...
    ]
  }
"""
import os, json, re
from docx import Document

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX = "/Users/javo/Downloads/Examenes_15x100_PNP_2026.docx"
OUT  = os.path.join(ROOT, "data", "simulacros.json")

LETRAS = "ABCDE"

def es_cabecera_materia(t):
    """Cabecera de materia: mayúsculas, no numerada, sin (A)/(B)."""
    if not t or len(t) < 15: return False
    if t.startswith(tuple(f"{L})" for L in LETRAS)): return False
    if re.match(r"^\d+\.", t): return False
    if "EXAMEN DE SIMULACRO" in t: return False
    if "CLAVE DE RESPUESTAS" in t or "RESPONDE EN HOJA" in t: return False
    # Debe estar mayormente en mayúsculas
    letras = [c for c in t if c.isalpha()]
    if not letras: return False
    up = sum(1 for c in letras if c.isupper()) / len(letras)
    return up > 0.85 and len(t) > 15

def parse_docx():
    doc = Document(DOCX)
    examenes = []
    current_exam = None
    current_materia = None
    current_preg = None
    en_clave = False

    for p in doc.paragraphs:
        t = p.text.strip()
        if not t: continue

        # Nuevo examen
        m = re.match(r"EXAMEN DE SIMULACRO N°\s*(\d+)", t)
        if m:
            current_exam = {
                "n": int(m.group(1)),
                "titulo": t,
                "preguntas": [],
            }
            examenes.append(current_exam)
            current_materia = None
            current_preg = None
            en_clave = False
            continue

        # Solo activar clave si la línea ES el título de la clave (mayúsculas, cortita)
        # NO la instrucción inicial "La clave de respuestas está al final..."
        if re.match(r"^\s*CLAVE DE RESPUESTAS", t) and len(t) < 60:
            en_clave = True
            continue

        if not current_exam or en_clave: continue

        # Cabecera de materia (línea antes de la 1ª pregunta o entre bloques)
        if es_cabecera_materia(t):
            current_materia = t.rstrip(".").strip()
            continue

        # Alternativa A)/B)/C)/D)/E)
        m = re.match(r"^([A-E])\)\s*(.+)$", t)
        if m and current_preg is not None:
            texto = m.group(2).rstrip(".").strip()
            current_preg["opciones"].append(texto)
            continue

        # Pregunta numerada
        m = re.match(r"^(\d+)\.\s+(.+)$", t)
        if m:
            current_preg = {
                "n": int(m.group(1)),
                "materia": current_materia or "",
                "pregunta": m.group(2).rstrip(".").rstrip(":").strip(),
                "opciones": [],
                "correcta": None,
            }
            current_exam["preguntas"].append(current_preg)
            continue

    # Ahora las tablas de claves (una por examen, en orden)
    for i, tabla in enumerate(doc.tables):
        if i >= len(examenes): break
        exam = examenes[i]
        claves = {}
        for row in tabla.rows[1:]:  # saltar header
            cells = [c.text.strip() for c in row.cells]
            # patrón: N | Resp | N | Resp | ... (5 pares)
            for j in range(0, len(cells), 2):
                if j+1 < len(cells):
                    num_txt = cells[j].strip()
                    resp_txt = cells[j+1].strip()
                    if num_txt.isdigit() and resp_txt in LETRAS:
                        claves[int(num_txt)] = LETRAS.index(resp_txt)
        # Asignar correcta a cada pregunta del examen
        for preg in exam["preguntas"]:
            if preg["n"] in claves:
                preg["correcta"] = claves[preg["n"]]

    return examenes

def main():
    examenes = parse_docx()
    total = sum(len(e["preguntas"]) for e in examenes)
    total_ok = sum(1 for e in examenes for p in e["preguntas"] if p["correcta"] is not None and len(p["opciones"])==5)
    total_5opts = sum(1 for e in examenes for p in e["preguntas"] if len(p["opciones"]) == 5)
    print(f"Exámenes: {len(examenes)}")
    print(f"Preguntas totales: {total}")
    print(f"Preguntas con 5 opciones: {total_5opts}")
    print(f"Preguntas con clave asignada: {total_ok}")

    # Reporte por examen
    for e in examenes:
        n_ok = sum(1 for p in e["preguntas"] if p["correcta"] is not None and len(p["opciones"])==5)
        n_bad = len(e["preguntas"]) - n_ok
        marca = "✓" if n_ok == 100 else f"⚠ {n_bad} incompletas"
        print(f"  Examen {e['n']:>2}: {len(e['preguntas']):>3} preguntas · {n_ok} completas {marca}")

    out = {
        "version": 1,
        "total_examenes": len(examenes),
        "total_preguntas": total,
        "materias_proporcion": {
            "Constitución": 7, "DUDH": 1, "DL 1267 (Ley PNP)": 6, "DL 1149 (Carrera)": 7,
            "DL 1291 (Anticorrupción)": 12, "DL 1318 (Formación)": 7, "Ley Ascensos": 2,
            "Ley 27444 (Proc. Admin.)": 45, "DS 009-2018": 2, "Ley 32130": 3,
            "DL 1611 (Prevención)": 7, "DL 1428 (Desaparecidos)": 1,
        },
        "examenes": examenes,
    }
    with open(OUT, "w") as f:
        json.dump(out, f, ensure_ascii=False)
    print(f"\n✓ Escrito {OUT}  ({os.path.getsize(OUT)//1024} KB)")

if __name__ == "__main__":
    main()
